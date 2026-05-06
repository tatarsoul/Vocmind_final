from __future__ import annotations

import io
from pathlib import Path
from typing import Any


def _import_docx_document():
    try:
        from docx import Document  # type: ignore
        return Document
    except Exception as exc:  # pragma: no cover - runtime dependency guard
        raise RuntimeError(
            "Для экспорта в DOCX не установлен пакет python-docx. Установите его командой: pip install python-docx"
        ) from exc


def _import_reportlab():
    try:
        from reportlab.lib.pagesizes import A4  # type: ignore
        from reportlab.lib.styles import getSampleStyleSheet  # type: ignore
        from reportlab.pdfbase import pdfmetrics  # type: ignore
        from reportlab.pdfbase.ttfonts import TTFont  # type: ignore
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # type: ignore
        return {
            "A4": A4,
            "getSampleStyleSheet": getSampleStyleSheet,
            "pdfmetrics": pdfmetrics,
            "TTFont": TTFont,
            "Paragraph": Paragraph,
            "SimpleDocTemplate": SimpleDocTemplate,
            "Spacer": Spacer,
        }
    except Exception as exc:  # pragma: no cover - runtime dependency guard
        raise RuntimeError(
            "Для экспорта в PDF не установлен пакет reportlab. Установите его командой: pip install reportlab"
        ) from exc


def export_runtime_status() -> dict[str, bool]:
    status = {"pdf": True, "docx": True}
    try:
        _import_docx_document()
    except RuntimeError:
        status["docx"] = False
    try:
        _import_reportlab()
    except RuntimeError:
        status["pdf"] = False
    return status


def _find_unicode_font() -> str | None:
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/Arial.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return str(path)
    return None


_FONT_REGISTERED = False


def _ensure_pdf_font(reportlab: dict[str, Any]) -> str:
    global _FONT_REGISTERED
    font_name = "Helvetica"
    font_path = _find_unicode_font()
    if font_path:
        font_name = "VocMindUnicode"
        if not _FONT_REGISTERED:
            reportlab["pdfmetrics"].registerFont(reportlab["TTFont"](font_name, font_path))
            _FONT_REGISTERED = True
    return font_name



def build_export_sections(meeting: dict[str, Any]) -> list[tuple[str, str]]:
    protocol_json = meeting.get("protocol_json") or {}
    analytics_json = meeting.get("analytics_json") or {}

    sections: list[tuple[str, str]] = []
    sections.append(("Общая информация", "\n".join([
        f"Название: {meeting.get('title') or 'Встреча'}",
        f"Статус: {meeting.get('status') or '—'}",
        f"Начало: {meeting.get('started_at') or '—'}",
        f"Окончание: {meeting.get('finished_at') or '—'}",
        f"Длительность, сек: {meeting.get('duration_seconds') or 0}",
        f"Режим: {meeting.get('capture_mode') or '—'}",
    ])))

    # Если у нас уже есть готовый protocol_text (его собирает
    # render_protocol_text — там уже все разделы: Кратко / Темы / Главные
    # тезисы / Решения / Задачи / Идеи / Риски / Следующие шаги), то выводим
    # ТОЛЬКО его. Отдельные секции из protocol_json ниже использовать НЕ нужно
    # — раньше это давало дубль (Кратко и Идеи рисовались дважды) и баг с
    # подменой topics → Идеи. Если по какой-то причине protocol_text пустой
    # (старые встречи в БД) — тогда падаем на старую логику.
    protocol_text_value = (meeting.get("protocol_text") or "").strip()
    if protocol_text_value:
        sections.append(("Протокол", protocol_text_value))
    elif protocol_json:
        # Fallback: protocol_text нет (например старая встреча) — рисуем
        # секции из JSON по отдельности. Здесь же фиксим старый баг
        # подмены topics → ideas: ideas рисуем только если они реально есть.
        if protocol_json.get("summary"):
            sections.append(("Кратко", str(protocol_json.get("summary") or "")))
        topics = [str(x).strip() for x in (protocol_json.get("topics") or []) if str(x).strip()]
        if topics:
            sections.append(("Темы", "\n".join(f"• {x}" for x in topics)))
        key_points = [str(x).strip() for x in (protocol_json.get("key_points") or []) if str(x).strip()]
        if key_points:
            sections.append(("Главные тезисы", "\n".join(f"• {x}" for x in key_points)))
        if protocol_json.get("decisions"):
            sections.append(("Решения", "\n".join(
                f"• {x}" for x in protocol_json.get("decisions") or [] if str(x).strip()
            )))
        if protocol_json.get("action_items"):
            action_lines = []
            for item in protocol_json.get("action_items") or []:
                if not isinstance(item, dict):
                    continue
                task = (item.get("task") or "").strip()
                if not task:
                    continue
                assignee = (item.get("assignee") or "").strip()
                due = (item.get("due") or "").strip()
                suffix = []
                if assignee:
                    suffix.append(f"ответственный: {assignee}")
                if due:
                    suffix.append(f"срок: {due}")
                action_lines.append(f"• {task}" + (f" ({'; '.join(suffix)})" if suffix else ""))
            if action_lines:
                sections.append(("Задачи", "\n".join(action_lines)))
        # ideas РОВНО ideas — никаких подмен на topics
        idea_items = [str(x).strip() for x in (protocol_json.get("ideas") or []) if str(x).strip()]
        if idea_items:
            sections.append(("Идеи", "\n".join(f"• {x}" for x in idea_items)))
        risk_items = [str(x).strip() for x in (protocol_json.get("risks") or []) if str(x).strip()]
        if risk_items:
            sections.append(("Риски", "\n".join(f"• {x}" for x in risk_items)))
        next_steps = [str(x).strip() for x in (protocol_json.get("next_steps") or []) if str(x).strip()]
        if next_steps:
            sections.append(("Следующие шаги", "\n".join(f"• {x}" for x in next_steps)))


    if analytics_json:
        balance = analytics_json.get("talk_balance") or {}
        analytics_lines = [
            f"Слов в разговоре: {analytics_json.get('word_count')}",
            f"Вопросов: {analytics_json.get('questions_count')}",
            f"Тем: {analytics_json.get('topics_count')}",
            f"Решений: {analytics_json.get('decisions_count')}",
            f"Задач: {analytics_json.get('action_items_count')}",
            f"Рисков: {analytics_json.get('risks_count')}",
            f"Баланс речи ME/THEM: {balance.get('me_percent')}% / {balance.get('them_percent')}%",
        ]
        keywords = analytics_json.get("top_keywords") or []
        if keywords:
            analytics_lines.append("Ключевые слова: " + ", ".join(map(str, keywords)))
        sections.append(("Аналитика", "\n".join(analytics_lines)))

    if meeting.get("transcript_text"):
        sections.append(("Транскрипт", meeting["transcript_text"]))

    return sections



def make_docx_bytes(meeting: dict[str, Any]) -> bytes:
    Document = _import_docx_document()
    doc = Document()
    doc.add_heading(meeting.get("title") or "VocMind — встреча", level=0)
    for title, body in build_export_sections(meeting):
        doc.add_heading(title, level=1)
        for chunk in str(body or "").split("\n"):
            doc.add_paragraph(chunk)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()



def make_pdf_bytes(meeting: dict[str, Any]) -> bytes:
    reportlab = _import_reportlab()
    bio = io.BytesIO()
    font_name = _ensure_pdf_font(reportlab)
    doc = reportlab["SimpleDocTemplate"](
        bio,
        pagesize=reportlab["A4"],
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
        title=meeting.get("title") or "VocMind export",
    )
    styles = reportlab["getSampleStyleSheet"]()
    styles["Normal"].fontName = font_name
    styles["Normal"].leading = 15
    styles["Title"].fontName = font_name
    styles["Heading1"].fontName = font_name
    story = [
        reportlab["Paragraph"]((meeting.get("title") or "VocMind — встреча").replace("\n", "<br/>"), styles["Title"]),
        reportlab["Spacer"](1, 12),
    ]
    for title, body in build_export_sections(meeting):
        story.append(reportlab["Paragraph"](title.replace("\n", "<br/>"), styles["Heading1"]))
        story.append(reportlab["Spacer"](1, 6))
        text = "<br/>".join(
            str(body or "")
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .split("\n")
        )
        story.append(reportlab["Paragraph"](text or "—", styles["Normal"]))
        story.append(reportlab["Spacer"](1, 12))
    doc.build(story)
    return bio.getvalue()
